#!/usr/bin/env python3
"""
Exporta registry.csv a un Excel de seguimiento de precios, ordenado por fecha.

Uso:
    python exportar_excel.py [--salida ARCHIVO.xlsx] [--desc] [--base DIR]

Columnas: Index | Glosa | N° OC | N° Cotización | Precio neto | Cantidad |
          Fecha cotización | Fecha OC

- Orden: por fecha de cotización (si falta, se usa la fecha de OC), de la más
  antigua a la más reciente. Con --desc, la más reciente primero.
- Las cotizaciones registradas solo por precios (nro_oc -1) aparecen como
  "SIN OC" y sin fecha de OC.
- Los precios del registro vienen en formatos mezclados ($8,720 / $8.720 /
  $7.133,41 / $953.54); aquí se convierten a número real.
"""

import argparse
import csv
import os
import re
from datetime import datetime, date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def a_numero(s):
    """Convierte un precio/cantidad en cualquiera de los formatos del registro."""
    s = re.sub(r"[^\d.,\-]", "", str(s or ""))
    if not s or not re.search(r"\d", s):
        return None
    if "," in s and "." in s:
        # el último separador es el decimal
        if s.rfind(",") > s.rfind("."):
            s = s.replace(".", "").replace(",", ".")
        else:
            s = s.replace(",", "")
    elif "," in s:
        dec = s.split(",")[-1]
        s = s.replace(",", "") if len(dec) == 3 and s.count(",") >= 1 else s.replace(",", ".")
    elif "." in s:
        dec = s.split(".")[-1]
        if len(dec) == 3 or s.count(".") > 1:
            s = s.replace(".", "")  # puntos de miles: $8.720 / $1.234.567
    try:
        n = float(s)
    except ValueError:
        return None
    return int(n) if n == int(n) else round(n, 2)


def a_fecha(s):
    s = str(s or "").strip()
    for fmt in ("%d-%m-%Y", "%d/%m/%Y", "%d-%m-%y", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def sin_oc(nro):
    try:
        return int(str(nro).strip()) <= 0
    except ValueError:
        return True


def main():
    ap = argparse.ArgumentParser(description="Exporta registry.csv a Excel ordenado por fecha.")
    ap.add_argument("--salida", help="Ruta del .xlsx (default: output/registro_precios_<fecha>.xlsx)")
    ap.add_argument("--desc", action="store_true", help="Más reciente primero")
    ap.add_argument("--base", default=os.path.dirname(os.path.abspath(__file__)))
    args = ap.parse_args()

    filas = []
    with open(os.path.join(args.base, "registry.csv"), newline="", encoding="utf-8") as f:
        for r in csv.DictReader(f):
            es_sin_oc = sin_oc(r.get("nro_oc"))
            f_cot = a_fecha(r.get("date_of_quotation"))
            f_oc_txt = "" if es_sin_oc else (r.get("fecha_oc") or "").strip()
            f_oc = a_fecha(f_oc_txt)
            idx = r.get("index", "")
            filas.append({
                "index": int(idx) if str(idx).isdigit() else idx,
                "glosa": (r.get("description") or "").strip(),
                "oc": "SIN OC" if es_sin_oc else int(r["nro_oc"]),
                "cot": (r.get("quotation_number") or "").strip(),
                "precio": a_numero(r.get("unit_price")),
                "cant": a_numero(r.get("quantity")),
                "f_cot": f_cot,
                "f_oc": f_oc if f_oc else (f_oc_txt or None),  # p.ej. "06-2025" queda como texto
                "orden": f_cot or f_oc or date.min,
            })

    filas.sort(key=lambda x: (x["orden"], str(x["index"])), reverse=args.desc)

    wb = Workbook()
    ws = wb.active
    ws.title = "Registro de precios"

    encabezados = ["Index", "Glosa", "N° OC", "N° Cotización", "Precio neto",
                   "Cantidad", "Fecha cotización", "Fecha OC"]
    anchos = [9, 55, 10, 16, 14, 11, 16, 13]

    fuente = Font(name="Arial", size=10)
    fuente_enc = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    relleno_enc = PatternFill("solid", start_color="3A5894")
    relleno_sin_oc = PatternFill("solid", start_color="F2F2F2")
    borde = Border(bottom=Side(style="thin", color="D9D9D9"))

    ws.append(encabezados)
    for c, ancho in enumerate(anchos, start=1):
        celda = ws.cell(row=1, column=c)
        celda.font = fuente_enc
        celda.fill = relleno_enc
        celda.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(c)].width = ancho
    ws.row_dimensions[1].height = 20

    for x in filas:
        ws.append([x["index"], x["glosa"], x["oc"], x["cot"], x["precio"],
                   x["cant"], x["f_cot"], x["f_oc"]])
        fila = ws.max_row
        for c in range(1, 9):
            celda = ws.cell(row=fila, column=c)
            celda.font = fuente
            celda.border = borde
            if x["oc"] == "SIN OC":
                celda.fill = relleno_sin_oc
        ws.cell(row=fila, column=3).alignment = Alignment(horizontal="center")
        ws.cell(row=fila, column=4).alignment = Alignment(horizontal="center")
        p = ws.cell(row=fila, column=5)
        p.number_format = '"$"#,##0.00' if isinstance(x["precio"], float) else '"$"#,##0'
        q = ws.cell(row=fila, column=6)
        q.number_format = "#,##0.##" if isinstance(x["cant"], float) else "#,##0"
        for c in (7, 8):
            celda = ws.cell(row=fila, column=c)
            celda.alignment = Alignment(horizontal="center")
            if isinstance(celda.value, date):
                celda.number_format = "dd-mm-yyyy"

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:H{ws.max_row}"

    # Nota al pie
    nota = ws.max_row + 2
    ws.cell(row=nota, column=1,
            value=(f"Generado el {date.today():%d-%m-%Y} desde registry.csv. "
                   "Precio neto = precio unitario sin IVA. "
                   "Filas en gris: cotizaciones registradas solo por precios (sin OC)."
                   )).font = Font(name="Arial", size=9, italic=True, color="666666")

    salida = args.salida or os.path.join(args.base, "output",
                                         f"registro_precios_{date.today():%Y-%m-%d}.xlsx")
    os.makedirs(os.path.dirname(os.path.abspath(salida)), exist_ok=True)
    wb.save(salida)

    sin_fecha = sum(1 for x in filas if x["f_cot"] is None)
    sin_precio = sum(1 for x in filas if x["precio"] is None)
    print(f"Excel generado: {salida}")
    print(f"  {len(filas)} filas | {sum(1 for x in filas if x['oc'] == 'SIN OC')} sin OC | "
          f"{sin_fecha} sin fecha de cotización | {sin_precio} sin precio legible")


if __name__ == "__main__":
    main()
